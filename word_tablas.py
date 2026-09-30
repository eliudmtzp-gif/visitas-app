import psycopg2
import os
from dotenv import load_dotenv
from docx import Document
from docx.shared import Inches

load_dotenv()
DATABASE_URL = os.getenv("DATABASE_URL")

def exportar_grupo_a_word(grupo_prefix, nombre_archivo):
    conn = psycopg2.connect(DATABASE_URL)
    cur = conn.cursor()

    # Consulta filtrada por grupo (ej. 'QA', 'QB')
    cur.execute("""
        SELECT identificador, nombre, direccion, telefono, variante,
               entrego, actualmente_la_visita, observaciones, fecha, ver_en_maps
        FROM visitas
        WHERE identificador LIKE %s
        ORDER BY identificador
    """, (f"{grupo_prefix}%",))

    registros = cur.fetchall()
    columnas = [desc[0] for desc in cur.description]

    # Crear documento Word
    doc = Document()
    doc.add_heading(f"Registros del grupo {grupo_prefix}", level=1)

    tabla = doc.add_table(rows=1, cols=len(columnas))
    tabla.style = 'Table Grid'

    # Encabezados
    encabezado = tabla.rows[0].cells
    for i, col in enumerate(columnas):
        encabezado[i].text = col.upper()

    # Filas
    for fila in registros:
        fila_tabla = tabla.add_row().cells
        for i, valor in enumerate(fila):
            fila_tabla[i].text = str(valor) if valor is not None else ""

    # Guardar documento
    doc.save(nombre_archivo)
    print(f"✅ Documento generado: {nombre_archivo}")

    cur.close()
    conn.close()

# Ejemplo de uso
exportar_grupo_a_word("PG", "grupo_PG.docx")
